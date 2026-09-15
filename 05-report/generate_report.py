#!/usr/bin/env python3
"""
Generates Legal-Ethical-Impact-Report.docx for Operation Hydra.
Requires: pip install python-docx
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x38, 0x64)
STEEL = RGBColor(0x44, 0x54, 0x6A)
ACCENT = RGBColor(0x2E, 0x74, 0xB5)

doc = Document()

# ---------- base style ----------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

for i, (size, color) in enumerate([(20, NAVY), (16, STEEL)], start=1):
    style = doc.styles[f"Heading {i}"]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color


def add_page_number_field(paragraph, kind="PAGE"):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = kind
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def body(text):
    doc.add_paragraph(text)


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def add_table(headers, rows, widths_in):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, htext in enumerate(headers):
        hdr_cells[i].text = htext
        for p_ in hdr_cells[i].paragraphs:
            for r in p_.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "1F3864")
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)
        table.columns[i].width = Inches(widths_in[i])
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
            table.columns[i].width = Inches(widths_in[i])
    doc.add_paragraph("")
    return table


# ================= COVER PAGE =================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for _ in range(6):
    doc.add_paragraph("")
r = title.add_run("OPERATION HYDRA")
r.font.size = Pt(36)
r.font.bold = True
r.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Legal, Ethical & Impact Report")
r.font.size = Pt(20)
r.font.bold = True
r.font.color.rgb = STEEL

tag = doc.add_paragraph()
tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tag.add_run("Unraveling a Multi-Vector Cyber Crime Involving Phishing, Spoofing & Financial Fraud")
r.italic = True
r.font.size = Pt(12)
r.font.color.rgb = STEEL

for _ in range(4):
    doc.add_paragraph("")

assign = doc.add_paragraph()
assign.alignment = WD_ALIGN_PARAGRAPH.CENTER
assign.add_run("Assignment 2 — Unit 2: Types of Cyber Crimes").font.size = Pt(12)

repo = doc.add_paragraph()
repo.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = repo.add_run("Repository: github.com/<your-username>/operation-hydra-forensics")
r.font.size = Pt(10)
r.font.color.rgb = ACCENT

doc.add_page_break()

# ================= HEADER / FOOTER =================
section = doc.sections[0]
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("Operation Hydra — Legal, Ethical & Impact Report")
hr.font.size = Pt(8)
hr.font.color.rgb = STEEL

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number_field(fp, "PAGE")
fp.add_run(" / ")
add_page_number_field(fp, "NUMPAGES")

# ================= BODY =================

h1("1. Executive Summary")
body("This report presents the legal, ethical, and stakeholder-impact analysis for \u201cOperation Hydra,\u201d "
     "a simulated multi-vector cyber crime campaign combining domain spoofing, phishing, malware-based "
     "credential theft, and financial fraud. The technical investigation \u2014 phishing header analysis, "
     "SPF/DKIM/DMARC evaluation, malware payload structure analysis, and financial fraud tracing \u2014 is "
     "documented in full in the accompanying repository. This document synthesizes those findings into a "
     "legal-technical brief: it maps each identified crime to applicable statutes, evaluates harm across "
     "five stakeholder groups, surfaces ethical tensions that arise during real investigations of this "
     "kind, and closes with policy and awareness recommendations.")
body("All technical artifacts referenced in this report \u2014 phishing samples, malware structure, "
     "transaction logs, and WHOIS records \u2014 are simulated for coursework purposes and do not reference "
     "real victims, real malicious infrastructure, or functional malicious code.")

h1("2. Cybercrime Breakdown & Legal Remedies")
body("Operation Hydra was decomposed into six distinct offenses along a single attack chain: domain "
     "spoofing enabling phishing, phishing enabling malware delivery, malware enabling credential/OTP "
     "theft, and stolen credentials enabling financial fraud that was then laundered through mule accounts "
     "and a simulated crypto off-ramp. Each stage constitutes a separate, prosecutable offense under Indian "
     "and international law, summarized below.")

h2("2.1 Crime-to-Statute Mapping")
add_table(
    ["Crime", "IT Act, 2000", "IPC, 1860 / BNS, 2023", "Global Equivalent"],
    [
        ["Email/Domain Spoofing", "\u00a766C, \u00a766D", "\u00a7463, \u00a7468, \u00a7471", "Budapest Art. 7; CFAA \u00a71030"],
        ["Phishing", "\u00a766D", "\u00a7420", "Budapest Art. 8"],
        ["Malware Distribution / Unauthorized Access", "\u00a743, \u00a766", "\u2014", "CFAA \u00a71030; Budapest Art. 2, 6"],
        ["Identity / Data Theft", "\u00a766C, \u00a772", "\u2014", "GDPR Art. 32/33; Budapest Art. 2"],
        ["Financial / Payment Fraud", "\u00a766D", "\u00a7420", "Budapest Art. 8; CFAA \u00a71030"],
        ["Money Laundering", "\u00a766, \u00a743 (instrumental)", "\u00a7120B", "Budapest Art. 25 (MLA); FATF standards"],
    ],
    [2.0, 1.3, 1.6, 1.9],
)

body("The full kill-chain narrative and section-by-section statutory justification are documented in "
     "01-classification/cybercrime-mapping.md. Two remedies deserve emphasis here. First, IT Act \u00a743A "
     "creates civil liability for a body corporate that fails to maintain reasonable security practices "
     "and thereby causes wrongful loss \u2014 relevant if the impersonated bank's own systems contributed to "
     "the campaign's plausibility. Second, IPC \u00a7120B (criminal conspiracy) is the charge most likely to "
     "bring in secondary participants such as mule-account holders and money launderers who may never have "
     "touched the phishing infrastructure directly.")

h2("2.2 Cross-Border Enforcement Gap")
body("A recurring practical problem in cases like Hydra is jurisdiction. The simulated header analysis "
     "placed originating infrastructure outside India (WHOIS registrant countries of Panama and China in "
     "the simulated records), which is realistic: real phishing infrastructure is routinely hosted or "
     "registered abroad specifically to complicate Indian law enforcement's ability to compel disclosure "
     "or extradite. The Budapest Convention on Cybercrime provides a mutual legal assistance framework for "
     "signatory states, but India is not a signatory, which narrows the available cross-border cooperation "
     "channels to bilateral treaties and informal law-enforcement-to-law-enforcement requests \u2014 both "
     "slower than an evidence trail that goes cold within days.")

h1("3. Impact on Stakeholders")
body("A single campaign like Hydra generates harm across several distinct groups, each experiencing a "
     "different kind of loss and requiring a different remedy.")

h2("3.1 Direct Victims (Account Holders)")
bullet("Financial: direct monetary loss (\u20b9138,000+ traced in the simulated fraud scenario), plus the "
       "burden of proving the transaction was unauthorized to recover funds under RBI's limited-liability circular.")
bullet("Emotional: victims of OTP-harvesting and account-takeover fraud commonly report acute anxiety and "
       "a lasting loss of trust in digital banking, independent of whether funds are eventually recovered.")
bullet("Time and administrative cost: filing police complaints, bank disputes, and credit monitoring after "
       "identity theft imposes a burden disproportionate to the amount stolen in many cases.")

h2("3.2 Impersonated Institutions (Bank / UPI Provider)")
bullet("Reputational: repeated impersonation erodes customer trust in the brand even though the "
       "institution's own systems were not breached.")
bullet("Financial: reimbursement obligations under RBI customer-liability rules, plus the cost of takedown "
       "requests, SPF/DKIM/DMARC hardening, and customer communication campaigns.")
bullet("Regulatory: obligation to report the incident and, if any EU data subjects were affected, GDPR "
       "breach-notification duties within 72 hours.")

h2("3.3 Unwitting Participants (Mule Account Holders)")
body("Individuals who allow their accounts to be used to receive and forward stolen funds \u2014 sometimes "
     "knowingly for a fee, sometimes as victims of a separate job-offer scam \u2014 occupy an ambiguous "
     "position. They face account freezes, potential criminal liability under \u00a7120B or as accessories, "
     "and reputational harm, even where their own culpability is unclear. This ambiguity is itself an "
     "impact worth naming: the fraud-tracing analysis in Section 4 of the technical repository identified "
     "a single mule account receiving funds from two separate victims, which is the kind of pattern "
     "investigators use to distinguish a knowing participant from an unwitting one.")

h2("3.4 Payment and Telecom Infrastructure Providers")
bullet("UPI switch operators and payment gateways bear compliance and monitoring costs to detect "
       "structuring patterns like the sub-\u20b950,000 transfers observed in the simulated trace.")
bullet("Telecom/SMS providers face pressure to filter OTP-relay abuse without introducing false positives "
       "that block legitimate one-time passwords.")

h2("3.5 Broader Public and Regulators")
body("Each successful campaign of this kind lowers aggregate trust in digital payment systems, which is a "
     "policy concern for regulators pursuing financial inclusion goals, and it consumes limited "
     "law-enforcement and CERT-In investigative capacity that would otherwise go to other cases.")

h1("4. Ethical Dilemmas")
body("Investigating and prosecuting campaigns like Hydra raises tensions that don't have clean answers "
     "even when the legal position is clear.")

h2("4.1 Crypto Tracing vs. Financial Privacy")
body("Tracing funds onto a blockchain off-ramp, as modeled in the fraud-tracing simulation, requires "
     "either voluntary exchange cooperation or compelled disclosure \u2014 both of which sit in tension with "
     "the pseudonymity that legitimate cryptocurrency users also rely on. A subpoena broad enough to "
     "identify a laundering wallet's real-world owner can also expose unrelated transaction history for "
     "that individual, and blockchain analytics firms that provide this tracing capability to law "
     "enforcement operate with limited public oversight of their methods and error rates. The ethical "
     "question is not whether tracing should happen, but how narrowly it can be scoped and how errors "
     "(false attribution) are corrected.")

h2("4.2 Deceptive Honeypots and Entrapment")
body("A common investigative technique against phishing infrastructure is a honeypot: a fake \u201cvictim\u201d "
     "account or credential set designed to lure the attacker into revealing more about their "
     "infrastructure or intent. This is broadly defensible when it passively observes an attacker who "
     "initiated contact, but the ethical and legal line blurs if investigators actively encourage an "
     "attacker toward conduct they would not otherwise have committed \u2014 the entrapment concern familiar "
     "from physical-world sting operations. A second, less-discussed dilemma is that honeypot "
     "infrastructure convincing enough to fool a sophisticated attacker is often equally convincing to a "
     "real, uninvolved user who stumbles onto it, which creates a duty of care investigators must design around.")

h2("4.3 Responsible Disclosure vs. Public Warning")
body("Once investigators identify active phishing domains and C2 infrastructure, there is a tension "
     "between disclosing details publicly (warning potential victims immediately) and keeping them "
     "confidential (preserving the investigation's ability to monitor the infrastructure and identify "
     "further victims or the operator before it is taken down). Neither choice is free of cost; the "
     "ethical weighting depends on the immediacy of ongoing harm to the public versus the investigative "
     "value of continued monitoring.")

h2("4.4 Liability Attribution Under Uncertainty")
body("Distinguishing a knowing money mule from a scam victim recruited under a fake job offer often "
     "cannot be done with certainty at the time accounts are frozen. Treating all mule accounts as equally "
     "culpable risks punishing people who were themselves defrauded; treating none as culpable risks "
     "giving genuine launderers a standing defense. This is less a technical problem than a due-process "
     "design problem: what evidentiary threshold should trigger criminal referral versus civil account "
     "freeze pending investigation.")

h1("5. Public Awareness & Policy Recommendations")

h2("5.1 Technical / Regulatory")
bullet("Mandate DMARC enforcement at p=reject (not p=none) for regulated financial institutions and their "
       "approved communication domains, closing the gap identified in Sample 2's analysis.")
bullet("Require payment switches to flag structuring patterns \u2014 multiple transfers just under a "
       "reporting threshold to the same or related payees within a short window \u2014 as a standard real-time "
       "rule, not a post-hoc audit finding.")
bullet("Extend mandatory attachment-type filtering (blocking double-extension executables such as "
       "*.pdf.exe) at the gateway level for institutions handling financial transactions.")

h2("5.2 Public Awareness")
bullet("Campaigns should specifically target OTP-harvesting patterns rather than generic \u201cdon't click "
       "suspicious links\u201d messaging \u2014 the technical analysis in this report found OTP-harvesting to be "
       "the highest-impact vector because it defeats two-factor authentication directly.")
bullet("Bank and UPI provider communications should consistently state a policy of never requesting OTP "
       "confirmation via inbound link, so any message that does so is self-evidently fraudulent regardless "
       "of how convincing its branding is.")

h2("5.3 Institutional / Legal")
bullet("Streamline the mule-account evidentiary process described in Section 4.4, distinguishing an "
       "interim civil freeze (low evidentiary bar, reversible) from a criminal referral (higher bar), so "
       "unwitting participants are not treated identically to knowing launderers by default.")
bullet("Pursue expanded bilateral mutual-legal-assistance arrangements given India's non-membership in the "
       "Budapest Convention, to reduce the cross-border evidence-gathering delay identified in Section 2.2.")

h1("6. Conclusion")
body("Operation Hydra illustrates how a single campaign chains together offenses that are individually "
     "well-covered by existing law \u2014 spoofing, phishing, unauthorized access, fraud, and laundering all "
     "map cleanly to specific IT Act, IPC, and international provisions \u2014 but whose investigation and "
     "remediation nonetheless surface hard, unresolved tensions: how narrowly to trace crypto off-ramps, "
     "how far honeypot deception can ethically go, when to warn the public versus keep monitoring quietly, "
     "and how to fairly sort victims from culpable participants among mule-account holders. Closing the "
     "technical gaps identified here (DMARC enforcement, structuring detection, attachment filtering) "
     "would blunt this specific attack pattern, but the ethical and jurisdictional questions in Sections 4 "
     "and 2.2 will recur under the next campaign's specific technical details even if this one is fully remediated.")

h1("References")
body("Information Technology Act, 2000 (as amended 2008).")
body("Indian Penal Code, 1860; Bharatiya Nyaya Sanhita, 2023.")
body("Council of Europe, Convention on Cybercrime (Budapest Convention), 2001.")
body("Computer Fraud and Abuse Act, 18 U.S.C. \u00a71030.")
body("General Data Protection Regulation (EU) 2016/679, Art. 32\u201333.")
body("Reserve Bank of India, Customer Protection \u2014 Limiting Liability of Customers in Unauthorised "
     "Electronic Banking Transactions, Circular, 2017.")
body("Full technical analysis, tooling, and simulated evidence: "
     "github.com/<your-username>/operation-hydra-forensics")

doc.save("Legal-Ethical-Impact-Report.docx")
print("Saved Legal-Ethical-Impact-Report.docx")